import glob
import io
import logging
import os
import pdf2image
import pptx

# constant for conversion from EMU to pixel
emu_as_pxl = 9525

# representation of pdf image to be added to a powerpoint presentation
class pdfImg:
	# default values
	height = 0
	left = 0
	top = 0
	width = 0

	# convert page in pdf document to png image, read width and height and generate bytes object
	def __init__(self, pdf, i):
		self.image = pdf2image.convert_from_path(pdf, first_page=i, last_page=i)[0]
		self.width = self.image.size[0] * emu_as_pxl
		self.height = self.image.size[1] * emu_as_pxl
		self.generate_bytes_obj()
	
	# adjust image height if necessary
	def adjust_height(self, avail_height, avail_width, total_image_height):
		adj_height = avail_height * self.height / total_image_height
		self.width = self.width * adj_height / self.height
		self.height = adj_height
		self.left = (avail_width - self.width) / 2
	
	# check if image width is bigger than slide width and adjust if necessary
	def adjust_width(self, avail_width):
		width_margin = avail_width - self.image.width
		if width_margin < 0:
			self.image.height = self.image.height * avail_width / self.image.width
			self.image.width = avail_width
		else:
			self.image.left = width_margin / 2
	
	# convert image to bytes object
	def generate_bytes_obj(self):
		bytes_obj = io.BytesIO()
		self.image.save(bytes_obj, "PNG")
		bytes_obj.seek(0)
		self.bytes = bytes_obj

# representation of powerpoint slide containing pdf images
class pptSlide:
	# default values
	height_margin = 0
	height_margin_fraction = 0
	total_height_images = 0

	# determine ppt slide object values
	def __init__(self, pdf, page_idx, ppt):
		self.width = ppt.slide_width
		self.height = ppt.slide_height
		# always set images to empty list to avoid previous images being stored
		self.images = []
		self.add_pdfImgs(pdf, page_idx)
		# determine total height of all images
		self.total_height_images = sum([image.height for image in self.images])
		# calculate y axis margin
		self.height_margin = self.height - self.total_height_images
		# adjust image height if total image height is bigger than slide
		if self.height_margin < 0:
			for image in self.images:
				image.adjust_height(self.height, self.width, self.total_height_images)
		else:
			# determine margin between images and slide border
			self.height_margin_fraction = self.height_margin / (len(self.images) + 1)
		self.determine_imgs_top_position()

	# add pdf images to ppt slide object
	def add_pdfImgs(self, pdf, page_idx):
		if not any(isinstance(i, int) for i in page_idx):
			logging.error("input needs to be list of integers")
			raise TypeError
		for i in page_idx:
			img = pdfImg(pdf, i)
			img.adjust_width(self.width)
			self.images.append(img)
	
	# determine how images should be positioned on slide's y axis
	def determine_imgs_top_position(self):
		position = self.height_margin_fraction
		for image in self.images:
			image.top = position
			position += image.height + self.height_margin_fraction
	
	# add pdf images to powerpoint slide
	def put_imgs_in_ppt(self, slide):
		for image in self.images:
			slide.shapes.add_picture(image.bytes, height=image.height, left=image.left, top=image.top, width=image.width)

	
# convert a page from a pdf document to image and append to powerpoint presentation
def add_pdf_pages_as_imgs_to_ppt(pdf, pdf_image_idx_per_slide, ppt):
	presentation = pptx.Presentation(ppt)
	for pdf_image_idx in pdf_image_idx_per_slide:
		slide = pptSlide(pdf, pdf_image_idx, presentation)
		slide.put_imgs_in_ppt(presentation.slides.add_slide(presentation.slide_layouts[6]))	
	presentation.save(ppt)
	return

# get tumor mutational burden label
def get_tmb_string(val):
	level = "Not available"
	if(0 <= val <= 5):
		level = "lav"
	elif(5 < val <= 20):
		level = "intermediær"
	elif(val > 20):
		level = "høy"
	return "{} mut/Mb; {}\n".format(val, level)

# use glob to find file in different folder structures
def glob_tsoppi_file(is_error, root, run_id, *path_units):
	glob_string_ous = os.path.join(root, '{}_TSO_500_LocalApp_postprocessing_results'.format(run_id), *path_units)
	glob_string_hus = os.path.join(root, run_id, '{}_ppr'.format(run_id), *path_units)
	files = []
	for gl in (glob_string_ous, glob_string_hus):
		files.extend(glob.glob(gl))
	if len(files) == 1:
		return files[0]
	elif not is_error:
		logging.warning("unsuccessful glob strings for {}:\n{}\n{}".format(run_id, glob_string_ous, glob_string_hus))
	else:
		logging.error("unsuccessful glob strings for {}:\n{}\n{}".format(run_id, glob_string_ous, glob_string_hus))
		raise ValueError
